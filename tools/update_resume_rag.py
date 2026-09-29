from copy import deepcopy
from pathlib import Path

from docx import Document


SOURCE = Path(r"D:\workspace\简历\last\王飞阳_算法工程师_正式版.docx")
OUTPUT = Path(r"D:\workspace\python\RAG_project\my_pro\artifacts\王飞阳_算法工程师_RAG修缮版.docx")


REPLACEMENTS = {
    "▪ 熟悉 Agent 任务规划、Tool Calling、结构化输出与长上下文处理；熟悉 RAG 全链路（文档解析与切块、稠密/稀疏混合检索、RRF 融合与重排、离线评测），具备 LangGraph/DeepAgents 多 Agent 编排与 RAG 应用落地经验。":
        "▪ 熟悉 Agent 任务规划、Tool Calling、结构化输出与上下文管理；具备 LangGraph/DeepAgents 多 Agent 编排及 RAG 应用落地经验。",
    "▪ 熟悉 MySQL、ChromaDB/Milvus 与 Pandas，具备向量检索、稠密/稀疏混合检索、CrossEncoder 重排与 RAGAS 评估实践。":
        "▪ 熟悉 MySQL、ChromaDB/Milvus、BGE-M3、RRF、CrossEncoder 与 RAGAS，具备父子块索引、稠密/稀疏混合检索、重排、评测与 Bad Case 分析实践。",
    "V93K 半导体测试机智能问答系统   算法工程师":
        "V93K 半导体测试知识智能问答系统   算法工程师",
    "整体流程：多格式资料解析与 OCR → 父子块切分与 BGE-M3 稠密/稀疏向量入库 → FAQ 前置匹配 → LLM 领域与策略路由 → 双路召回与 RRF 融合 → 父块还原与 Rerank 精排 → 带来源约束的生成与拒答 → SSE 流式返回并落库多轮会话。":
        "项目简介：参与三人团队面向封测车间测试工程师建设可溯源 RAG 问答服务，整合 V93K 手册、常见错误、STIL 语法与测试 API；本人重点负责父子块索引、混合检索、策略路由及多轮记忆模块。",
    "技术栈：Python + FastAPI/Uvicorn + LangChain + BGE-M3 / BGE-Reranker-Large + ChromaDB + GLM-4.5-Air + MySQL + RapidOCR + RAGAS + SSE":
        "技术栈：Python + FastAPI/Uvicorn + LangChain + BGE-M3 / BGE-Reranker-Large + ChromaDB + GLM-4.5-Air + MySQL + RapidOCR + RAGAS",
    "▪ 资料解析与索引：实现 PDF/Word/PPT/图片/Markdown 多格式解析与 RapidOCR 兜底，按父块 1200、子块 300、重叠 50 字符递归切分；子块写入 ChromaDB 并记录 parent_id、领域与来源元数据，父块原文单独存储，供检索后还原完整上下文。":
        "▪ 检索架构与父子索引：负责父子块索引及混合检索链路，按父块 1200、子块 300、重叠 50 字符切分；基于 BGE-M3 构建稠密/稀疏双路召回与 RRF 融合，按 parent_id 还原父块后用 BGE-Reranker-Large 精排 Top-3，兼顾语义问法及错误码、API 名等精确词项。",
    "▪ 混合检索与精排：用 BGE-M3 同时生成 1024 维稠密向量与稀疏词项权重，稠密路走 ChromaDB、稀疏路走内存词项索引，按 RRF 融合排名；子块召回 Top-5 后按 parent_id 还原父块，再用 BGE-Reranker-Large 对「原问题 × 父块」精排取 Top-3，兼顾语义问法与错误码、API 名等精确词项。":
        "▪ 策略路由与可信回答：负责 FAQ 前置匹配（阈值 0.85）及 Direct/HyDE/SubQuery/Backtracking 策略路由，服务端执行枚举校验并兜底 Direct；通过来源约束、证据不足拒答和来源回传控制幻觉。",
    "▪ 路由与检索策略：FAQ 前置以标准问题向量的余弦相似度（阈值 0.85）命中即返回标准答案，跳过路由、检索、重排与生成；未命中时由一次 LLM 调用输出知识/闲聊判定、五领域（工艺整合/设备维护/良率分析/洁净室/制程质量）与四种检索策略（Direct / HyDE / SubQuery / Backtracking），服务端做枚举校验并兜底 direct。":
        "▪ 多轮记忆模块：负责基于 MySQL 持久化会话记录与摘要，生成时组合会话摘要和最近 3 轮历史；通过摘要压缩早期对话、保留近期原文，在控制上下文长度的同时支持追问与上下文衔接。",
    "▪ 系统在封测车间内部试运行，知识库覆盖 V93K 手册、常见错误、STIL 语法与测试 API 四类资料，切分为 25 个父块、120 个子块，支持按领域与来源过滤。":
        "▪ 内部试运行覆盖 V93K 手册、常见错误、STIL 语法与测试 API 四类资料，形成 25 个父块、120 个子块，并支持按领域和来源过滤。",
    "▪ FAQ 前置直通让高频标准问题跳过路由、检索、重排与生成全链路，命中时无需调用大模型即可返回答案，降低平均响应时延与算力消耗。":
        "▪ FAQ 前置直通使高频标准问题跳过路由、检索、重排和生成，命中后无需调用大模型即可返回标准答案，降低平均响应时延与算力消耗。",
    "▪ SSE 流式输出把首字可见时间缩短到秒级；direct 链路在冒烟样本上单题（检索 + 重排 + 生成）约 6~9 秒，并按题留痕，用于分阶段区分解析、召回、融合、重排与生成的耗时。":
        "▪ 多轮记忆采用“会话摘要 + 最近 3 轮历史”组合控制上下文长度，支持追问衔接；Direct 链路在冒烟样本上单题（检索 + 重排 + 生成）约 6～9 秒，并按题留痕，支持分阶段定位召回、融合、重排与生成耗时。",
}


REMOVE_PARAGRAPHS = {
    "▪ 接口与多轮会话：FastAPI 提供 JSON 与 SSE 双通道接口（meta / chunk / done 事件），MySQL 持久化会话记录与摘要，生成时注入摘要与最近 3 轮历史；Prompt 约束只依据检索上下文作答，证据不足时明确拒答，并回传来源文档名供人工核对。",
    "▪ 评测与复盘：构建覆盖 4 类知识源的 25 条人工 QA 数据集，在固定 direct 口径下用 RAGAS 统计 Context Precision / Recall，逐题记录上下文、耗时与异常样本；据此定位出「judge 失败被 NaN 跳过、导致综合分虚高」和「RRF 未对双路重合项累加」两类缺陷，并给出指标有效率与四组消融的改进方案。",
}


def replace_paragraph_text(paragraph, new_text):
    """Replace visible text while preserving paragraph and first-run formatting."""
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = new_text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(new_text)


def copy_run_format(source_run, target_run):
    if source_run._element.rPr is not None:
        target_run._element.insert(0, deepcopy(source_run._element.rPr))


def normalize_bullet_paragraph(paragraph, bullet_template, body_template):
    text = paragraph.text
    body = text[2:] if text.startswith("▪ ") else text.lstrip("▪").lstrip()
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    bullet_run = paragraph.add_run("▪ ")
    body_run = paragraph.add_run(body)
    copy_run_format(bullet_template, bullet_run)
    copy_run_format(body_template, body_run)


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def main():
    doc = Document(SOURCE)
    matched = set()
    for paragraph in doc.paragraphs:
        old_text = paragraph.text
        if old_text in REPLACEMENTS:
            replace_paragraph_text(paragraph, REPLACEMENTS[old_text])
            matched.add(old_text)

    missing = set(REPLACEMENTS) - matched
    if missing:
        raise RuntimeError("Unmatched source paragraphs:\n" + "\n".join(sorted(missing)))

    removed = set()
    for paragraph in list(doc.paragraphs):
        if paragraph.text in REMOVE_PARAGRAPHS:
            removed.add(paragraph.text)
            remove_paragraph(paragraph)
    missing_removals = REMOVE_PARAGRAPHS - removed
    if missing_removals:
        raise RuntimeError("Unmatched removal paragraphs:\n" + "\n".join(sorted(missing_removals)))

    # Match the established resume style: blue bullet marker, dark-gray body text.
    bullet_template = doc.paragraphs[19].runs[0]
    body_template = doc.paragraphs[19].runs[1]
    rag_bullets = {
        value for value in REPLACEMENTS.values() if value.startswith("▪ ")
    }
    for paragraph in doc.paragraphs:
        if paragraph.text in rag_bullets:
            normalize_bullet_paragraph(paragraph, bullet_template, body_template)

    # Avoid splitting the next project's three-item responsibility list across pages.
    next_project = "封测产线报错日志根因分析系统   算法工程师"
    paragraphs = doc.paragraphs
    project_index = next(i for i, p in enumerate(paragraphs) if p.text == next_project)
    responsibility_index = next(
        i for i in range(project_index + 1, len(paragraphs))
        if paragraphs[i].text == "个人职责"
    )
    for index in (responsibility_index, responsibility_index + 1, responsibility_index + 2):
        paragraphs[index].paragraph_format.keep_with_next = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    print(f"updated_paragraphs={len(matched)}")


if __name__ == "__main__":
    main()
